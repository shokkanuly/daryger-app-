import re
from docx import Document
from typing import List, Dict
from .ocr import extract_prices_from_line, clean_name

def get_cell_text_clean(cell) -> str:
    text_elements = cell._element.xpath(".//w:t[not(ancestor::w:del)]")
    return "".join(node.text for node in text_elements if node.text)

def parse_docx(file_path: str) -> List[Dict]:
    doc = Document(file_path)
    rows = []
    
    for table in doc.tables:
        for row in table.rows:
            if not row.cells:
                continue
            cleaned_row = [get_cell_text_clean(cell).strip() for cell in row.cells]
            if not any(cleaned_row):
                continue
            
            name_cell = ""
            prices = []
            for cell in cleaned_row:
                cell_clean = re.sub(r'[^\d.,]', '', cell).replace(",", ".")
                if cell_clean:
                    parts = cell_clean.split(".")
                    if 3 <= len(parts[0]) <= 7:
                        try:
                            prices.append(float(cell_clean))
                            continue
                        except ValueError:
                            pass
                if cell and len(cell) > len(name_cell):
                    name_cell = cell

            name_cell = clean_name(name_cell).strip(" -:.\t")
            if name_cell and len(name_cell) > 4 and prices:
                resident = prices[0]
                nonresident = prices[1] if len(prices) >= 2 else resident
                rows.append({
                    "name": name_cell,
                    "price_resident": resident,
                    "price_nonresident": nonresident,
                    "currency": "KZT"
                })

    if not rows:
        # fallback to paragraph lines
        for para in doc.paragraphs:
            text_elements = para._element.xpath(".//w:t[not(ancestor::w:del)]")
            text = "".join(node.text for node in text_elements if node.text)
            
            name, res, nonres = extract_prices_from_line(text)
            if name and res:
                rows.append({
                    "name": name,
                    "price_resident": res,
                    "price_nonresident": nonres,
                    "currency": "KZT"
                })

    return rows
