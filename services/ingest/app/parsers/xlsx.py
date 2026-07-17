import re
import openpyxl
from typing import List, Dict
from .ocr import clean_name

def detect_columns(row_values: List[str]):
    name_col = -1
    price_res_col = -1
    price_nonres_col = -1

    for idx, val in enumerate(row_values):
        val_lower = str(val).lower()
        if any(kw in val_lower for kw in ["услуг", "наимен", "исслед", "анализ", "name", "service", "description"]):
            name_col = idx
        elif any(kw in val_lower for kw in ["нерезидент", "non-resident", "nonresident"]) and any(kw in val_lower for kw in ["цена", "стоим", "price", "cost", "₸", "тг"]):
            price_nonres_col = idx
        elif any(kw in val_lower for kw in ["резидент", "resident"]) and any(kw in val_lower for kw in ["цена", "стоим", "price", "cost", "₸", "тг"]):
            price_res_col = idx
        elif any(kw in val_lower for kw in ["цена", "стоим", "стоиомость", "price", "cost", "₸", "тг", "сумма"]) and price_res_col == -1:
            price_res_col = idx

    return name_col, price_res_col, price_nonres_col

def parse_xlsx(file_path: str) -> List[Dict]:
    wb = openpyxl.load_workbook(file_path, data_only=True)
    rows = []

    for sheet in wb.worksheets:
        header_row_idx = -1
        name_col = -1
        price_res_col = -1
        price_nonres_col = -1

        # Scan the first 30 rows to find header
        max_scan = min(30, sheet.max_row)
        for r_idx in range(1, max_scan + 1):
            row_vals = [sheet.cell(r_idx, c_idx).value for c_idx in range(1, sheet.max_column + 1)]
            row_vals_str = [str(v) if v is not None else "" for v in row_vals]
            
            n_col, pr_col, pnr_col = detect_columns(row_vals_str)
            if n_col != -1 and pr_col != -1:
                header_row_idx = r_idx
                name_col = n_col
                price_res_col = pr_col
                price_nonres_col = pnr_col
                break

        # Heuristic auto-profiling fallback if header detection failed
        if header_row_idx == -1 or name_col == -1 or price_res_col == -1:
            col_types = {c: {"num_count": 0, "text_len_sum": 0, "empty_count": 0, "total": 0} for c in range(sheet.max_column)}
            scan_limit = min(50, sheet.max_row)
            
            for r_idx in range(1, scan_limit + 1):
                for c_idx in range(1, sheet.max_column + 1):
                    val = sheet.cell(r_idx, c_idx).value
                    c_key = c_idx - 1
                    if val is None:
                        col_types[c_key]["empty_count"] += 1
                    else:
                        val_str = str(val).strip()
                        col_types[c_key]["total"] += 1
                        clean_num = re.sub(r'[^\d\.]', '', val_str.replace(",", "."))
                        if clean_num and len(clean_num) < 10:
                            try:
                                float(clean_num)
                                col_types[c_key]["num_count"] += 1
                            except ValueError:
                                pass
                        col_types[c_key]["text_len_sum"] += len(val_str)

            best_name_col = -1
            max_text_len = -1
            for c, stats in col_types.items():
                if stats["total"] > 0:
                    avg_len = stats["text_len_sum"] / stats["total"]
                    num_ratio = stats["num_count"] / stats["total"]
                    if num_ratio < 0.3 and avg_len > 12 and stats["text_len_sum"] > max_text_len:
                        max_text_len = stats["text_len_sum"]
                        best_name_col = c

            best_price_col = -1
            max_num_count = -1
            for c, stats in col_types.items():
                if c != best_name_col and stats["total"] > 0:
                    num_ratio = stats["num_count"] / stats["total"]
                    if num_ratio > 0.4 and stats["num_count"] > max_num_count:
                        max_num_count = stats["num_count"]
                        best_price_col = c

            if best_name_col != -1:
                name_col = best_name_col
            else:
                name_col = 0

            if best_price_col != -1:
                price_res_col = best_price_col
            else:
                price_res_col = 1
                
            start_row = 1
        else:
            start_row = header_row_idx + 1

        # Extract data starting from the row after the header
        for r_idx in range(start_row, sheet.max_row + 1):
            name_val = sheet.cell(r_idx, name_col + 1).value
            price_res_val = sheet.cell(r_idx, price_res_col + 1).value
            
            if price_nonres_col != -1:
                price_nonres_val = sheet.cell(r_idx, price_nonres_col + 1).value
            else:
                price_nonres_val = price_res_val

            if name_val is None:
                continue

            name_str = clean_name(str(name_val)).strip(" -:.\t")
            if len(name_str) < 4:
                continue

            # Parse and validate prices
            try:
                res_clean = str(price_res_val).strip() if price_res_val is not None else ""
                nonres_clean = str(price_nonres_val).strip() if price_nonres_val is not None else ""

                # Avoid parsing descriptions as price (too many letters)
                if len(re.findall(r'[a-zA-Zа-яА-ЯёЁ]', res_clean)) > 4:
                    continue

                res_str = re.sub(r'[^\d\.]', '', res_clean.replace(",", "."))
                nonres_str = re.sub(r'[^\d\.]', '', nonres_clean.replace(",", "."))
                
                res_price = float(res_str) if res_str else None
                nonres_price = float(nonres_str) if nonres_str else res_price

                if res_price is not None and res_price > 0:
                    rows.append({
                        "name": name_str,
                        "price_resident": res_price,
                        "price_nonresident": nonres_price if nonres_price is not None else res_price,
                        "currency": "KZT"
                    })
            except Exception:
                pass
                
    wb.close()
    return rows
