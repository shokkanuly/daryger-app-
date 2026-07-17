import json
import os
import openpyxl
import re

def main():
    json_path = "prisma/seed-data/services.json"
    xlsx_path = "prisma/seed-data/Справочник услуг.xlsx"
    
    # 1. Load existing services
    existing_services = []
    existing_names = set()
    if os.path.exists(json_path):
        with open(json_path, "r", encoding="utf-8") as f:
            existing_services = json.load(f)
            for s in existing_services:
                existing_names.add(s["name"].lower().strip())
                for syn in s.get("synonyms", []):
                    existing_names.add(syn.lower().strip())

    print(f"Loaded {len(existing_services)} existing services from JSON.")

    # 2. Parse Excel
    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    sheet = wb.active
    
    # Header mapping
    header = [cell.value for cell in sheet[1]]
    col_idx = {name: i for i, name in enumerate(header)}
    
    new_services_count = 0
    
    for r_idx in range(2, sheet.max_row + 1):
        name_ru = sheet.cell(r_idx, col_idx.get("Name_ru") + 1).value
        if not name_ru or str(name_ru).strip().lower() in ("none", ""):
            continue
            
        name_ru = str(name_ru).strip()
        name_key = name_ru.lower()
        
        if name_key in existing_names:
            continue
            
        code = str(sheet.cell(r_idx, col_idx.get("Code") + 1).value or "").strip()
        specialty = str(sheet.cell(r_idx, col_idx.get("Специальность") + 1).value or "").strip()
        
        # Categorize based on keywords
        name_lower = name_ru.lower()
        if any(w in name_lower for w in ["анализ", "определение", "исследование крови", "исследование мочи", "пцр", "днк", "рнк", "сыворотк"]):
            category = "lab"
        elif any(w in name_lower for w in ["прием", "консультация", "осмотр врача"]):
            category = "consult"
        elif any(w in name_lower for w in ["узи", "ультразвуков", "мрт", "кт", "рентген", "томограф"]):
            category = "diagnostic"
        else:
            category = "procedure"
            
        # Build synonyms
        synonyms = [name_ru]
        
        # Add abbreviation synonyms automatically
        if "аланинаминотрансфераза" in name_lower:
            synonyms.append("АЛТ")
        if "аспартатаминотрансфераза" in name_lower:
            synonyms.append("АСТ")
        if "тиреотропный" in name_lower:
            synonyms.append("ТТГ")
        if "электрокардио" in name_lower:
            synonyms.append("ЭКГ")
        if "фиброгастро" in name_lower or "гастроскопия" in name_lower:
            synonyms.append("ФГДС")
        if "магнитно-резонансная томография" in name_lower:
            synonyms.append("МРТ")
        if "компьютерная томография" in name_lower:
            synonyms.append("КТ")
        if "ультразвуковое исследование" in name_lower:
            synonyms.append("УЗИ")
            # Extract organs if possible
            match_org = re.search(r'исследование\s+([а-яА-ЯёЁ\s]+)', name_lower)
            if match_org:
                org_name = match_org.group(1).strip()
                synonyms.append(f"УЗИ {org_name}")

        existing_services.append({
            "name": name_ru,
            "synonyms": list(set(synonyms)),
            "category": category,
            "icdCode": code if code and code.lower() != "none" else None
        })
        
        existing_names.add(name_key)
        new_services_count += 1

    # Save back to services.json
    with open(json_path, "w", encoding="utf-8") as f:
        json.dump(existing_services, f, ensure_ascii=False, indent=2)
        
    print(f"Added {new_services_count} new services from Excel. Total standard services: {len(existing_services)}.")

if __name__ == "__main__":
    main()
